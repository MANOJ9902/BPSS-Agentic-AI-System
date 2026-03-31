
import os
import sys
import json
import time
import subprocess
from pathlib import Path
from datetime import date, datetime
from typing import Optional

import pandas as pd
from pypdf import PdfReader

# ── New google-genai SDK ─────────────────────────────────────────────────────
from google import genai
from google.genai import types

# ─────────────────────────────────────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────────────────────────────────────

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

DATA_DIR   = Path(os.environ.get("BPSS_DATA_DIR", "./data"))
MODEL      = "gemini-2.5-flash"
MAX_TOKENS = 4096

# Analyst review dates per candidate (from dataset)
REVIEW_DATES: dict[str, date] = {
    "CAND-101": date(2026, 2, 11),
    "CAND-102": date(2026, 2,  4),
    "CAND-103": date(2026, 2, 25),
    "CAND-104": date(2026, 2, 20),
    "CAND-105": date(2026, 3,  3),
    "CAND-106": date(2026, 1, 31),
}

# ─────────────────────────────────────────────────────────────────────────────
# Tool 1 — read_document
# ─────────────────────────────────────────────────────────────────────────────

def read_document(filename: str) -> str:
    """Return the full plain text of a PDF or DOCX file."""
    path = DATA_DIR / filename
    if not path.exists():
        available = [f.name for f in DATA_DIR.iterdir() if f.is_file()]
        return (
            f"ERROR: '{filename}' not found in {DATA_DIR}.\n"
            f"Available files: {', '.join(available)}"
        )

    ext = path.suffix.lower()

    if ext == ".pdf":
        try:
            reader = PdfReader(str(path))
            pages  = [p.extract_text() or "" for p in reader.pages]
            text   = "\n\n".join(pages).strip()
            return f"[SOURCE: {filename}]\n\n{text}"
        except Exception as e:
            return f"ERROR reading PDF '{filename}': {e}"

    if ext in (".docx", ".doc"):
        try:
            result = subprocess.run(
                ["pandoc", str(path), "-t", "plain", "--wrap=none"],
                capture_output=True, text=True, timeout=30,
            )
            if result.returncode == 0 and result.stdout.strip():
                return f"[SOURCE: {filename}]\n\n{result.stdout.strip()}"
        except FileNotFoundError:
            pass
        except Exception:
            pass

        try:
            from docx import Document
            doc  = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return f"[SOURCE: {filename}]\n\n{text}"
        except Exception as e:
            return f"ERROR reading DOCX '{filename}': {e}"

    return f"ERROR: Unsupported file type '{ext}' for '{filename}'"


# ─────────────────────────────────────────────────────────────────────────────
# Tool 2 — query_structured_data
# ─────────────────────────────────────────────────────────────────────────────

def query_structured_data(
    filename: str,
    sheet: Optional[str]      = None,
    filter_col: Optional[str] = None,
    filter_val: Optional[str] = None,
) -> str:
    """Load a CSV or XLSX file and return all (or filtered) rows as JSON."""
    path = DATA_DIR / filename
    if not path.exists():
        return f"ERROR: '{filename}' not found in {DATA_DIR}"

    ext = path.suffix.lower()
    try:
        if ext == ".csv":
            df = pd.read_csv(str(path))
        elif ext in (".xlsx", ".xlsm", ".xls"):
            try:
                df = pd.read_excel(str(path), sheet_name=sheet or 0)
            except Exception:
                import openpyxl
                wb    = openpyxl.load_workbook(str(path), read_only=True)
                names = wb.sheetnames
                return (
                    f"ERROR: Could not read sheet '{sheet}' from '{filename}'. "
                    f"Available sheets: {names}"
                )
        else:
            return f"ERROR: Unsupported extension '{ext}'"
    except Exception as e:
        return f"ERROR loading '{filename}': {e}"

    original_rows = len(df)

    if filter_col and filter_val:
        col = filter_col.strip()
        val = filter_val.strip()
        if col not in df.columns:
            return (
                f"ERROR: Column '{col}' not found in '{filename}'. "
                f"Columns: {list(df.columns)}"
            )
        df = df[df[col].astype(str).str.strip().str.upper() == val.upper()]

    source_tag = filename + (f" / sheet='{sheet}'" if sheet else "")
    meta = (
        f"[SOURCE: {source_tag} | "
        f"total rows: {original_rows} | "
        f"rows returned: {len(df)}"
        + (f" | filter: {filter_col}={filter_val}" if filter_col else "")
        + "]\n\n"
    )
    return meta + df.to_json(orient="records", indent=2, date_format="iso")


# ─────────────────────────────────────────────────────────────────────────────
# Tool 3 — check_document_freshness
# ─────────────────────────────────────────────────────────────────────────────

def check_document_freshness(candidate_id: str) -> str:
    """Check whether a candidate's address-proof documents are within 90 days."""
    cid = candidate_id.strip().upper()
    review = REVIEW_DATES.get(cid)
    if not review:
        return (
            f"ERROR: No review date on record for '{cid}'. "
            f"Known IDs: {list(REVIEW_DATES.keys())}"
        )

    path = DATA_DIR / "document_inventory.csv"
    if not path.exists():
        return "ERROR: document_inventory.csv not found in data directory"

    df    = pd.read_csv(str(path))
    rows  = df[df["candidate_id"].str.strip().str.upper() == cid]

    lines = [
        f"[SOURCE: document_inventory.csv | candidate: {cid} | review date: {review}]",
        "",
    ]

    addr_mask = rows["doc_type"].str.contains(
        r"Address|Statement|Utility|Bill|Council Tax|Proof",
        case=False, na=False,
    )
    addr_rows = rows[addr_mask]

    if addr_rows.empty:
        lines.append("❌  No address-proof document found in inventory.")
    else:
        for _, row in addr_rows.iterrows():
            present = str(row.get("present_in_folder", "?")).strip().lower()
            if present == "no":
                lines.append(
                    f"❌  {row['doc_type']} ({row['document_id']}): "
                    "document NOT retained in folder — cannot be audited."
                )
                continue

            raw_date = row.get("document_date")
            if pd.isna(raw_date) or str(raw_date).strip() in ("", "nan", "NaT", "None"):
                lines.append(
                    f"❌  {row['doc_type']} ({row['document_id']}): "
                    "date unknown — freshness cannot be assessed."
                )
                continue

            doc_date = datetime.strptime(str(raw_date)[:10], "%Y-%m-%d").date()
            age_days = (review - doc_date).days
            ok       = age_days <= 90
            symbol   = "✅" if ok else "❌"
            verdict  = "WITHIN" if ok else "EXCEEDS"
            lines.append(
                f"{symbol}  {row['doc_type']} ({row['document_id']}): "
                f"dated {doc_date}, {age_days} days before analyst review "
                f"— {verdict} 90-day limit."
            )

    return "\n".join(lines)


# ─────────────────────────────────────────────────────────────────────────────
# Tool 4 — list_files
# ─────────────────────────────────────────────────────────────────────────────

def list_files() -> str:
    """Return a manifest of all files in the data directory."""
    if not DATA_DIR.exists():
        return f"ERROR: data directory '{DATA_DIR}' does not exist."
    files = sorted(f for f in DATA_DIR.iterdir() if f.is_file())
    if not files:
        return f"ERROR: data directory '{DATA_DIR}' is empty."
    lines = [f"[SOURCE: directory listing of {DATA_DIR}]\n"]
    for f in files:
        lines.append(f"  {f.name:<55} {f.suffix.upper():<6}  {f.stat().st_size // 1024} KB")
    return "\n".join(lines)


# ─────────────────────────────────────────────────────────────────────────────
# Tool dispatcher
# ─────────────────────────────────────────────────────────────────────────────

def dispatch_tool(name: str, inputs: dict) -> str:
    if name == "read_document":
        return read_document(**inputs)
    if name == "query_structured_data":
        return query_structured_data(**inputs)
    if name == "check_document_freshness":
        return check_document_freshness(**inputs)
    if name == "list_files":
        return list_files()
    return f"ERROR: Unknown tool '{name}'"


# ─────────────────────────────────────────────────────────────────────────────
# Tool schemas (new google.genai format)
# ─────────────────────────────────────────────────────────────────────────────

TOOLS = [
    types.Tool(
        function_declarations=[
            types.FunctionDeclaration(
                name="read_document",
                description=(
                    "Extract and return the full text of a PDF or DOCX file from the data directory. "
                    "Use for policy documents, adjudication register, candidate packs, analyst notes, "
                    "and email approvals."
                ),
                parameters=types.Schema(
                    type=types.Type.OBJECT,
                    properties={
                        "filename": types.Schema(
                            type=types.Type.STRING,
                            description="Exact filename including extension. E.g. 'CAND-104_candidate_pack.docx'",
                        )
                    },
                    required=["filename"],
                ),
            ),
            types.FunctionDeclaration(
                name="query_structured_data",
                description=(
                    "Load a CSV or XLSX file and return rows as JSON with optional filtering. "
                    "Files: bpps_tracker_export.csv, BPSS_case_tracker.xlsx (sheets: tracker, "
                    "document_inventory, employment_history, risk_rules), document_inventory.csv, "
                    "employment_history.csv."
                ),
                parameters=types.Schema(
                    type=types.Type.OBJECT,
                    properties={
                        "filename":   types.Schema(type=types.Type.STRING, description="Filename including extension."),
                        "sheet":      types.Schema(type=types.Type.STRING, description="Sheet name for XLSX files."),
                        "filter_col": types.Schema(type=types.Type.STRING, description="Column name to filter on."),
                        "filter_val": types.Schema(type=types.Type.STRING, description="Value to match."),
                    },
                    required=["filename"],
                ),
            ),
            types.FunctionDeclaration(
                name="check_document_freshness",
                description=(
                    "Check whether a candidate's address-proof documents fall within "
                    "the 90-day policy window relative to the analyst review date."
                ),
                parameters=types.Schema(
                    type=types.Type.OBJECT,
                    properties={
                        "candidate_id": types.Schema(type=types.Type.STRING, description="E.g. 'CAND-102'")
                    },
                    required=["candidate_id"],
                ),
            ),
            types.FunctionDeclaration(
                name="list_files",
                description="Return a manifest of all data files available in the data directory.",
                parameters=types.Schema(type=types.Type.OBJECT, properties={}),
            ),
        ]
    )
]

# ─────────────────────────────────────────────────────────────────────────────
# System prompt
# ─────────────────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """\
You are a BPSS Screening Compliance Analyst AI. You reason carefully over
multi-format screening case files to answer compliance and risk questions.

POLICY RULES YOU MUST APPLY
1. A case may only be closed as Clear when ALL mandatory controls are complete
   OR when a formally approved exception exists in the Adjudication Register.
2. Mandatory controls for all standard hires: identity, right-to-work (RTW),
   employment history (3 years), criminality/basic disclosure.
3. Criminality is NOT mandatory for roles coded INTERN or CONTRACTOR-LTD.
4. Identity requires: at least one government-issued photo ID PLUS proof of
   address dated within 90 days of the analyst review date.
5. Expired BRP alone does NOT satisfy right-to-work evidence.
6. Employment history requires documentary or referee evidence covering the
   prior 3 years. CV only is WEAK evidence. Unexplained gaps >31 days must
   be accounted for.
7. ONLY entries in the Adjudication Register are approved exceptions.
   Manager emails, analyst notes, or verbal confirmation are NOT approved
   exceptions and do NOT satisfy mandatory controls.
8. "Ready to Join" in the tracker means business onboarding readiness ONLY.
   It is NOT equivalent to BPSS closure.
9. Tracker status fields are operational and may lag actual evidence.
   Always verify tracker claims against the underlying evidence documents.
10. If evidence is absent, unretained, or unverifiable, say so explicitly —
    do not infer it exists.

HOW TO ANSWER
- Always cite source files, sheet names, and specific fields/rows.
- When tracker data conflicts with document evidence, flag it as a
  contradiction and explain which source is authoritative.
- Never fabricate evidence. If something cannot be confirmed, say so.
- Structure answers clearly: conclusion first, then supporting evidence.\
"""


# ─────────────────────────────────────────────────────────────────────────────
# BPSSAgent
# ─────────────────────────────────────────────────────────────────────────────

class BPSSAgent:
    """
    Agentic AI that reasons over BPSS screening files using the new google.genai SDK.
    Maintains session history so follow-up questions work naturally.
    """

    def __init__(self, verbose: bool = True):
        self._client  = genai.Client(api_key=GEMINI_API_KEY)
        self.verbose  = verbose
        self._history: list[types.Content] = []

    def _log(self, msg: str) -> None:
        if self.verbose:
            print(msg, flush=True)

    def _separator(self, char: str = "─", width: int = 70) -> str:
        return char * width

    def _send_with_retry(self, contents, max_retries: int = 5):
        """Send a request with exponential backoff on rate-limit errors."""
        for attempt in range(max_retries):
            try:
                return self._client.models.generate_content(
                    model=MODEL,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=SYSTEM_PROMPT,
                        tools=TOOLS,
                        max_output_tokens=MAX_TOKENS,
                    ),
                )
            except Exception as e:
                err_str = str(e).lower()
                if any(k in err_str for k in ("resource", "exhausted", "429", "quota", "rate")):
                    wait = min(30 * (2 ** attempt), 90)
                    self._log(f"  ⏳  Rate-limited (attempt {attempt+1}/{max_retries}). Waiting {wait}s...")
                    time.sleep(wait)
                else:
                    raise
        return self._client.models.generate_content(
            model=MODEL,
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=SYSTEM_PROMPT,
                tools=TOOLS,
                max_output_tokens=MAX_TOKENS,
            ),
        )

    def ask(self, question: str) -> str:
        """Ask a question, maintaining session history. Returns the final answer."""
        self._log(f"\n{self._separator('=')}")
        self._log(f"QUESTION: {question}")
        self._log(self._separator('='))

        # Append the user turn to history
        self._history.append(
            types.Content(role="user", parts=[types.Part(text=question)])
        )

        max_iterations = 25

        for iteration in range(1, max_iterations + 1):

            response = self._send_with_retry(self._history)
            candidate = response.candidates[0]

            # Append model response to history
            self._history.append(candidate.content)

            # Collect function calls
            function_calls = [
                part.function_call
                for part in candidate.content.parts
                if part.function_call is not None
            ]

            # No tool calls → extract and return final answer
            if not function_calls:
                answer = "\n".join(
                    part.text
                    for part in candidate.content.parts
                    if hasattr(part, "text") and part.text
                ).strip()

                self._log(f"\n{self._separator()}")
                self._log("ANSWER:\n")
                self._log(answer)
                self._log(self._separator())
                return answer

            # Execute tools and build function-response turn
            response_parts = []
            for fc in function_calls:
                inputs = dict(fc.args) if fc.args else {}

                self._log(f"\n  🔧  [{iteration}] {fc.name}")
                self._log(f"       inputs: {json.dumps(inputs)}")

                result  = dispatch_tool(fc.name, inputs)
                preview = result[:300].replace("\n", " ")
                dots    = "..." if len(result) > 300 else ""
                self._log(f"       result: {preview}{dots}")

                response_parts.append(
                    types.Part(
                        function_response=types.FunctionResponse(
                            name=fc.name,
                            response={"result": result},
                        )
                    )
                )

            # Append tool results to history and loop
            self._history.append(
                types.Content(role="tool", parts=response_parts)
            )

        return "ERROR: Agent exceeded maximum iterations without producing a final answer."

    def reset(self) -> None:
        """Clear session history."""
        self._history = []
        self._log("Session history cleared.")


# ─────────────────────────────────────────────────────────────────────────────
# Batch runner
# ─────────────────────────────────────────────────────────────────────────────

EVALUATION_QUESTIONS: list[str] = [
    "Which candidate files are not yet ready for BPSS closure, and why?",
    "For each candidate, determine whether the identity, employment history, "
    "criminality, and right-to-work checks are complete. Cite evidence.",
    "Which cases appear to have policy exceptions? Distinguish approved "
    "exceptions from unapproved deviations.",
    "Compare the adjudication note for CAND-104 with the policy. "
    "Is closure justified?",
    "Which candidates have potentially stale documents or evidence that "
    "expired before review completion?",
    "Draft an escalation summary for the Screening Operations Lead covering "
    "the top 3 highest-risk cases.",
    "What information is missing to determine whether CAND-105 can be cleared?",
    "Reconcile any contradictions between the tracker and the analyst notes.",
    "Which candidates were marked 'ready to join' even though one or more "
    "mandatory controls were incomplete?",
    (
        "Create a structured JSON summary for each candidate with: "
        "status, control completion by category, missing items, "
        "risk level, and supporting evidence."
    ),
]


def run_batch(output_file: str = "bpss_answers.md") -> None:
    """Answer all 10 evaluation questions and write to a markdown file."""
    print(f"\nRunning batch — {len(EVALUATION_QUESTIONS)} questions")
    print(f"Output: {output_file}\n")

    lines = [
        "# BPSS Agent — Evaluation Answers\n",
        f"_Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}_\n\n",
        "---\n",
    ]

    for i, question in enumerate(EVALUATION_QUESTIONS, 1):
        if i > 1:
            print("⏳ Waiting 45s to avoid free-tier rate limits...")
            time.sleep(45)

        print(f"\n{'─'*60}")
        print(f"Q{i}/{len(EVALUATION_QUESTIONS)}: {question[:80]}…")

        agent = BPSSAgent(verbose=True)
        try:
            answer = agent.ask(question)
        except Exception as e:
            print(f"Error on Q{i}: {e}")
            answer = f"ERROR: {e}"

        lines.append(f"## Q{i}. {question}\n\n{answer}\n\n---\n")

    with open(output_file, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))

    print(f"\n\n✅  All answers written to: {output_file}")


# ─────────────────────────────────────────────────────────────────────────────
# Interactive CLI
# ─────────────────────────────────────────────────────────────────────────────

def run_cli() -> None:
    """Interactive terminal chat. Maintains session history for follow-ups."""
    agent = BPSSAgent(verbose=True)

    print("\n" + "═" * 60)
    print("  BPSS Agentic AI — Interactive Mode")
    print("  Commands: 'reset' | 'batch' | 'exit'")
    print("═" * 60 + "\n")

    while True:
        try:
            question = input("❓  Question: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n\nExiting.")
            break

        if not question:
            continue

        cmd = question.lower()
        if cmd in ("exit", "quit"):
            print("Exiting.")
            break
        if cmd == "reset":
            agent.reset()
            continue
        if cmd == "batch":
            run_batch()
            continue

        agent.ask(question)


# ─────────────────────────────────────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    args = sys.argv[1:]

    if not args:
        run_cli()

    elif args[0] == "--batch":
        out = args[1] if len(args) > 1 else "bpss_answers.md"
        run_batch(output_file=out)

    elif args[0] == "--ask" and len(args) > 1:
        question = " ".join(args[1:])
        agent    = BPSSAgent(verbose=True)
        agent.ask(question)

    else:
        print(
            "Usage:\n"
            "  python agent.py                         # interactive mode\n"
            "  python agent.py --batch                 # answer all 10 questions\n"
            "  python agent.py --batch out.md          # write to custom file\n"
            '  python agent.py --ask "your question"   # single question\n'
        )
